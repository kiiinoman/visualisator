"""
Сопоставление хостов и подсетей.

Логика:
  - Каждый хост из host.xlsx сопоставляется с подсетью из subnet.xlsx
    (ищется наиболее специфичная подсеть, в которую входит IP хоста)

  - unique_hosts   = хосты, для которых НЕ найдена ни одна подсеть в subnet.xlsx
  - unique_subnets = подсети из subnet.xlsx, в которые НЕ попал ни один хост

Входные файлы:
    subnet.xlsx  — колонка 'subnet'
    host.xlsx    — колонка 'host'

Результат — result.xlsx:
    Лист 'subnet_hosts'   — subnet | hosts | кол-во
    Лист 'unique_hosts'   — хосты без подсети
    Лист 'unique_subnets' — подсети без хостов

Python 3.9+  |  pip install pandas openpyxl
"""

import ipaddress
import logging
import sys
import os
from datetime import datetime
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── Файлы ─────────────────────────────────────────────────────────────────────
SUBNET_FILE = "subnet.xlsx"
HOST_FILE   = "host.xlsx"
OUTPUT_FILE = "result.xlsx"
LOG_FILE    = "run.log"
# ─────────────────────────────────────────────────────────────────────────────

C_HDR_GREEN = "2E7D32"
C_HDR_BLUE  = "1565C0"
C_ROW_EVEN  = "F5F5F5"
C_ROW_ODD   = "FFFFFF"
C_HAS_HOSTS = "C8E6C9"
C_NO_HOSTS  = "FFCDD2"


# ── Логгер ────────────────────────────────────────────────────────────────────
def setup_logger():
    log = logging.getLogger("make_result")
    log.setLevel(logging.DEBUG)
    fmt = logging.Formatter("%(asctime)s  %(levelname)-8s  %(message)s",
                            datefmt="%Y-%m-%d %H:%M:%S")
    fh = logging.FileHandler(LOG_FILE, encoding="utf-8", mode="w")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(fmt)
    ch = logging.StreamHandler(sys.stdout)
    ch.setLevel(logging.INFO)
    ch.setFormatter(fmt)
    log.addHandler(fh)
    log.addHandler(ch)
    return log


# ── Стили ─────────────────────────────────────────────────────────────────────
def _border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def hdr(cell, text, bg=C_HDR_GREEN):
    cell.value = text
    cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _border()

def dat(cell, value=None, bg=C_ROW_ODD, bold=False):
    if value is not None:
        cell.value = value
    cell.font = Font(name="Arial", size=10, bold=bold)
    cell.fill = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    cell.border = _border()


# ── Утилиты ───────────────────────────────────────────────────────────────────
def read_col(filepath, colname, log):
    # type: (str, str, logging.Logger) -> List[str]
    if not os.path.exists(filepath):
        log.error("Файл не найден: %s" % filepath)
        raise FileNotFoundError("Файл не найден: %s" % filepath)
    df = pd.read_excel(filepath, dtype=str)
    df.columns = df.columns.str.strip()
    log.debug("Колонки в %s: %s" % (filepath, list(df.columns)))
    mapping = {c.lower(): c for c in df.columns}
    real = mapping.get(colname.lower())
    if not real:
        raise ValueError("Колонка '%s' не найдена в %s. Есть: %s" % (
            colname, filepath, list(df.columns)))
    values = df[real].dropna().str.strip().tolist()
    log.info("  %s — прочитано строк: %d" % (filepath, len(values)))
    return values

def try_network(s):
    # type: (str) -> Optional[ipaddress.IPv4Network]
    try:
        return ipaddress.ip_network(s, strict=False)
    except ValueError:
        return None

def try_ip(s):
    # type: (str) -> Optional[ipaddress.IPv4Address]
    try:
        return ipaddress.ip_address(s)
    except ValueError:
        return None


# ── Главная логика ────────────────────────────────────────────────────────────
def main():
    log = setup_logger()
    log.info("=" * 60)
    log.info("Старт  %s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    log.info("=" * 60)

    # ── 1. Загружаем подсети ──────────────────────────────────────────────────
    log.info("ШАГ 1/4  Загрузка подсетей из %s" % SUBNET_FILE)
    raw_subnets = read_col(SUBNET_FILE, "subnet", log)

    # Нормализуем и дедуплицируем
    seen_sn = set()
    subnets_ordered = []   # List[str]  нормализованные CIDR, уникальные
    networks_ordered = []  # List[IPv4Network]  в том же порядке
    skipped_bad = 0
    skipped_dup = 0

    for s in raw_subnets:
        net = try_network(s)
        if net is None:
            log.warning("Нераспознанная подсеть пропущена: '%s'" % s)
            skipped_bad += 1
            continue
        key = str(net)
        if key in seen_sn:
            log.debug("Дубликат подсети пропущен: %s" % key)
            skipped_dup += 1
            continue
        seen_sn.add(key)
        subnets_ordered.append(key)
        networks_ordered.append(net)

    log.info("  Строк raw:              %d" % len(raw_subnets))
    log.info("  Нераспознанных:         %d" % skipped_bad)
    log.info("  Дублей удалено:         %d" % skipped_dup)
    log.info("  Уникальных подсетей:    %d" % len(subnets_ordered))

    # ── 2. Загружаем хосты ───────────────────────────────────────────────────
    log.info("ШАГ 2/4  Загрузка хостов из %s" % HOST_FILE)
    raw_hosts = read_col(HOST_FILE, "host", log)

    seen_h = set()
    hosts = []
    for h in raw_hosts:
        if h not in seen_h:
            seen_h.add(h)
            hosts.append(h)

    log.info("  Строк raw:              %d" % len(raw_hosts))
    log.info("  Дублей удалено:         %d" % (len(raw_hosts) - len(hosts)))
    log.info("  Уникальных хостов:      %d" % len(hosts))

    # ── 3. Сопоставление ─────────────────────────────────────────────────────
    log.info("ШАГ 3/4  Сопоставление хостов -> подсети")

    # subnet_str -> [host, ...]
    subnet_to_hosts = {s: [] for s in subnets_ordered}  # type: Dict[str, List[str]]
    unique_hosts = []

    for host in hosts:
        ip = try_ip(host)
        if ip is None:
            log.warning("Невалидный IP пропущен: '%s'" % host)
            unique_hosts.append(host)
            continue

        best_net = None
        best_prefix = -1
        for net, cidr in zip(networks_ordered, subnets_ordered):
            if ip in net and net.prefixlen > best_prefix:
                best_net = cidr
                best_prefix = net.prefixlen

        if best_net:
            subnet_to_hosts[best_net].append(host)
            log.debug("  %s  ->  %s" % (host, best_net))
        else:
            unique_hosts.append(host)
            log.debug("  %s  ->  [нет подсети]" % host)

    # Подсети с хостами и без
    matched_subnets = [s for s in subnets_ordered if subnet_to_hosts[s]]
    unique_subnets  = [s for s in subnets_ordered if not subnet_to_hosts[s]]

    log.info("  Хостов сопоставлено:    %d / %d" % (len(hosts) - len(unique_hosts), len(hosts)))
    log.info("  Хостов без подсети:     %d  -> лист unique_hosts" % len(unique_hosts))
    log.info("  Подсетей с хостами:     %d" % len(matched_subnets))
    log.info("  Подсетей без хостов:    %d  -> лист unique_subnets" % len(unique_subnets))

    # Топ-10 подсетей
    top10 = sorted(matched_subnets, key=lambda s: len(subnet_to_hosts[s]), reverse=True)[:10]
    log.info("  Топ-10 подсетей по кол-ву хостов:")
    for s in top10:
        log.info("    %-25s  %d хостов" % (s, len(subnet_to_hosts[s])))

    # ── 4. Запись ─────────────────────────────────────────────────────────────
    log.info("ШАГ 4/4  Запись %s" % OUTPUT_FILE)
    wb = Workbook()

    # Лист 1 — subnet_hosts (все подсети из subnet.xlsx)
    ws1 = wb.active
    ws1.title = "subnet_hosts"
    hdr(ws1.cell(1, 1), "subnet",        C_HDR_GREEN)
    hdr(ws1.cell(1, 2), "hosts",         C_HDR_GREEN)
    hdr(ws1.cell(1, 3), "кол-во хостов", C_HDR_GREEN)
    ws1.column_dimensions["A"].width = 22
    ws1.column_dimensions["B"].width = 60
    ws1.column_dimensions["C"].width = 16
    ws1.row_dimensions[1].height = 22
    for row, s in enumerate(subnets_ordered, start=2):
        hl = subnet_to_hosts[s]
        bg = C_HAS_HOSTS if hl else C_NO_HOSTS
        dat(ws1.cell(row, 1), s,                              bg=bg)
        dat(ws1.cell(row, 2), ", ".join(hl) if hl else "",    bg=bg)
        dat(ws1.cell(row, 3), len(hl),                        bg=bg, bold=bool(hl))
    log.info("  Лист 'subnet_hosts'   записан: %d строк" % len(subnets_ordered))

    # Лист 2 — unique_hosts
    ws2 = wb.create_sheet("unique_hosts")
    hdr(ws2.cell(1, 1), "host (нет в subnet.xlsx)", C_HDR_BLUE)
    ws2.column_dimensions["A"].width = 25
    ws2.row_dimensions[1].height = 22
    for i, h in enumerate(unique_hosts, start=2):
        dat(ws2.cell(i, 1), h, bg=C_ROW_EVEN if i % 2 == 0 else C_ROW_ODD)
    log.info("  Лист 'unique_hosts'   записан: %d строк" % len(unique_hosts))

    # Лист 3 — unique_subnets
    ws3 = wb.create_sheet("unique_subnets")
    hdr(ws3.cell(1, 1), "subnet (нет ни одного хоста)", C_HDR_BLUE)
    ws3.column_dimensions["A"].width = 25
    ws3.row_dimensions[1].height = 22
    for i, s in enumerate(unique_subnets, start=2):
        dat(ws3.cell(i, 1), s, bg=C_ROW_EVEN if i % 2 == 0 else C_ROW_ODD)
    log.info("  Лист 'unique_subnets' записан: %d строк" % len(unique_subnets))

    wb.save(OUTPUT_FILE)
    log.info("Файл сохранён: %s" % OUTPUT_FILE)
    log.info("=" * 60)
    log.info("Завершено успешно")
    log.info("=" * 60)

    print("\n✓ Готово! %s" % OUTPUT_FILE)
    print("  subnet_hosts   — %d подсетей (%d с хостами)" % (len(subnets_ordered), len(matched_subnets)))
    print("  unique_hosts   — %d хостов без подсети" % len(unique_hosts))
    print("  unique_subnets — %d подсетей без хостов" % len(unique_subnets))
    print("  Лог:             %s" % LOG_FILE)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        logging.getLogger("make_result").exception("КРИТИЧЕСКАЯ ОШИБКА: %s" % e)
        sys.exit(1)
